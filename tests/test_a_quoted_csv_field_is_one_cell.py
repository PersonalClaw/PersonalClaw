"""CSV text handed to ``sheet_create`` is read as CSV: a quoted field is one cell.

A field in quotes may hold the delimiter, a line break and doubled quotes, and it is still one
cell. ``sheet_create``'s ``csv`` text is read the way Knowledge reads a ``.csv`` file, by one
reading both share (``knowledge.readers.delimited_rows``), so a spreadsheet made from CSV text
has the cells the text has. A space after a comma is how people and models write CSV by hand,
so it does not stop the quoted field after it from being one cell.
"""

from __future__ import annotations

import io

import pytest

from personalclaw.knowledge.readers import FileReader


def _prov(tmp_path, monkeypatch):
    monkeypatch.setenv("PERSONALCLAW_HOME", str(tmp_path))
    from personalclaw.artifacts.native import NativeArtifactProvider

    return NativeArtifactProvider(root=tmp_path / "artifacts")


def _sheet_from(tmp_path, monkeypatch, text: str) -> list[list[object]]:
    """The cells of the spreadsheet ``sheet_create`` makes from *text*, as the file holds them."""
    from openpyxl import load_workbook

    from personalclaw.mcp_artifacts import _document_create

    prov = _prov(tmp_path, monkeypatch)
    reply = _document_create(prov, "sheet_create", {"name": "Contacts", "csv": text}, None, _quiet)
    assert "Error" not in reply, reply
    [made] = prov.list()
    body = prov.raw_bytes(made.slug)
    assert body is not None
    sheet = load_workbook(io.BytesIO(body[0])).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_a_quoted_field_holding_a_comma_is_one_cell(tmp_path, monkeypatch):
    text = 'Name,City\nAda,"Paris, France"\nGrace,Lisbon\n'

    assert _sheet_from(tmp_path, monkeypatch, text) == [
        ["Name", "City"],
        ["Ada", "Paris, France"],
        ["Grace", "Lisbon"],
    ]


def test_a_quoted_field_holding_a_line_break_is_one_cell(tmp_path, monkeypatch):
    text = 'Name,Notes\nAda,"first line\nsecond line"\nGrace,done\n'

    assert _sheet_from(tmp_path, monkeypatch, text) == [
        ["Name", "Notes"],
        ["Ada", "first line\nsecond line"],
        ["Grace", "done"],
    ]


def test_a_quoted_field_keeps_its_doubled_quotes_as_one_quote(tmp_path, monkeypatch):
    text = 'Name,Said\nAda,"she said ""yes"", twice"\n'

    assert _sheet_from(tmp_path, monkeypatch, text) == [
        ["Name", "Said"],
        ["Ada", 'she said "yes", twice'],
    ]


def test_a_quoted_field_after_a_comma_and_a_space_is_one_cell(tmp_path, monkeypatch):
    text = 'Name, City\nAda, "Paris, France"\n'

    assert _sheet_from(tmp_path, monkeypatch, text) == [
        ["Name", "City"],
        ["Ada", "Paris, France"],
    ]


def test_plain_csv_text_is_read_as_before(tmp_path, monkeypatch):
    """Unquoted cells, a blank line between rows and Windows line endings read as they did."""
    text = "Region,Q1\r\nEMEA,120\r\n\r\nAPAC,99.5\r\n"

    assert _sheet_from(tmp_path, monkeypatch, text) == [
        ["Region", "Q1"],
        ["EMEA", "120"],
        ["APAC", "99.5"],
    ]


def test_csv_text_it_cannot_read_makes_nothing_and_says_why(tmp_path, monkeypatch):
    from personalclaw.mcp_artifacts import _document_create

    prov = _prov(tmp_path, monkeypatch)
    audited: list[tuple[str, str, str]] = []

    def _audit(outcome, slug="", error=""):
        audited.append((outcome, slug, error))

    oversized_field = "x" * 200_000
    reply = _document_create(
        prov,
        "sheet_create",
        {"name": "Contacts", "csv": f"Name\n{oversized_field}\n"},
        None,
        _audit,
    )

    assert reply.startswith("Error: could not read the csv text:"), reply
    assert "field larger than field limit" in reply
    assert prov.list() == []
    assert audited[-1][0] == "denied"


@pytest.mark.parametrize("suffix, delimiter", [(".csv", ","), (".tsv", "\t")])
def test_knowledge_reads_a_quoted_field_after_a_space_as_one_cell(tmp_path, suffix, delimiter):
    """The reading ``sheet_create`` shares: a Knowledge file keeps the same cells."""
    path = tmp_path / f"contacts{suffix}"
    path.write_text(f'Name{delimiter} City\nAda{delimiter} "Paris, France"\n', encoding="utf-8")

    text, meta = FileReader().read(str(path))

    assert meta["row_count"] == 2
    assert text.splitlines() == ["| Name | City |", "| --- | --- |", "| Ada | Paris, France |"]


def _quiet(outcome, slug="", error=""):
    """The tool's audit callback; these tests read what was stored, not the audit row."""
