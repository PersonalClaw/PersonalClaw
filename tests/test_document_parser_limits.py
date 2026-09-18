""".docx/.xlsx/.pptx untrusted-input caps (#2747).

**The defect.** ``documents/``'s parsers hand raw bytes to python-docx / openpyxl /
python-pptx, each of which opens the archive itself, so none of ``doc_parser.py``'s shipped
zip-bomb posture (``_MAX_ZIP_ENTRY``, ``_MAX_DECOMPRESS``, ``_read_zip_entry``'s
actual-size check) reached them and nothing bounded the parse. Measured on ``main``: a
**56 KB** ``.docx`` holding 200,000 one-character paragraphs cost **153 s** of CPU, and at
that compression ratio the 16 MiB write cap admits roughly 12 hours and 23 GB.

**Two caps, refusing differently, and the difference is the point.**

* An ARCHIVE cap refuses: :class:`DocumentTooLarge` → ``document_too_large`` 413. There is
  no honest partial answer for a half-opened ZIP.
* A STRUCTURAL cap truncates and appends a ``size_limit`` loss. A raise would tell a user
  their 40,000-paragraph document is broken; the loss report says "too big, here is what
  fit" through the machinery the module already has for saying what it dropped.

**Calibration is a first-class assertion here**, not a comment. ``TestKnownCleanCorpus``
below is what stops a future tightening from refusing ordinary documents — an ordinary
2,000 x 100 spreadsheet carries 10.9 MB of XML in one part, which a cap chosen from the
"largest real file" alone (438 KB) would have refused outright.
"""

from __future__ import annotations

import ast
import io
import zipfile
from pathlib import Path

import pytest

from personalclaw.documents import limits
from personalclaw.documents.limits import (
    MAX_ARCHIVE_ENTRIES,
    MAX_BLOCKS,
    MAX_CELLS,
    MAX_SLIDES,
    MAX_XML_PART_BYTES,
    DocumentTooLarge,
    assert_archive_within_limits,
)

_FIXTURES = Path(__file__).parent / "fixtures"


# ── helpers: build the issue's own repro ─────────────────────────────────────


def _minimal_docx() -> bytes:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("x")
    doc.save(buf)
    return buf.getvalue()


def _docx_with(paragraphs: int) -> bytes:
    """The issue's repro verbatim: ``paragraphs`` one-character paragraphs in one part."""
    original = _minimal_docx()
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        body = archive.read("word/document.xml").decode()
    fat = body.replace("<w:p>", "<w:p><w:r><w:t>A</w:t></w:r></w:p>" * paragraphs + "<w:p>", 1)
    return _repack(original, {"word/document.xml": fat.encode()})


def _repack(original: bytes, replacements: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            dst.writestr(item.filename, replacements.get(item.filename, src.read(item.filename)))
    return out.getvalue()


def _oversized_package() -> bytes:
    """A well-formed ZIP whose one XML part decompresses past the per-part cap.

    Deliberately NOT a valid office document: the archive cap must refuse it BEFORE any
    document library is handed the bytes, so a parser that forgot the cap fails this with
    its library's own error rather than :class:`DocumentTooLarge` — which is exactly the
    distinction the assertion checks.
    """
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        dst.writestr("[Content_Types].xml", b"<a>" + b"x" * (MAX_XML_PART_BYTES + 1) + b"</a>")
    return out.getvalue()


# ── archive caps: a refusal, before the library ──────────────────────────────


class TestArchiveCaps:
    def test_an_oversized_xml_part_is_refused_and_names_the_part(self):
        with pytest.raises(DocumentTooLarge) as excinfo:
            assert_archive_within_limits(_oversized_package())
        assert "[Content_Types].xml" in str(excinfo.value)
        assert "50 MB" in str(excinfo.value)

    def test_the_check_never_consults_the_declared_size(self):
        """``ZipInfo.file_size`` is the header's CLAIM and a hostile archive is free to lie
        in it — the reason ``doc_parser._read_zip_entry`` reads bytes rather than trusting
        the declaration, and the posture this module was told to reuse.

        Asserted on the SOURCE rather than by forging a central directory: the property is
        "this code does not read the declaration", and a forged-header fixture proves only
        that one forgery is caught. The name appearing anywhere in the enforcement path is
        the regression worth catching.
        """
        source = Path(limits.__file__).read_text(encoding="utf-8")
        code = "\n".join(
            line for line in source.splitlines() if not line.strip().startswith("#")
        ).split('"""')[-1]
        assert "file_size" not in code, "the cap must measure decompressed bytes, not the header"

    def test_too_many_parts_is_refused(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for index in range(MAX_ARCHIVE_ENTRIES + 1):
                dst.writestr(f"part{index}.xml", b"<a/>")
        with pytest.raises(DocumentTooLarge, match="parts"):
            assert_archive_within_limits(out.getvalue())

    def test_media_is_not_charged_to_the_xml_budget(self):
        """A legitimate photo deck is mostly media, and charging it the XML budget would
        refuse the ordinary case the cap was written to allow."""
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as dst:
            dst.writestr("ppt/media/image1.png", b"\x00" * (MAX_XML_PART_BYTES + 1))
            dst.writestr("[Content_Types].xml", b"<a/>")
        assert_archive_within_limits(out.getvalue())  # must not raise

    def test_a_non_zip_is_left_to_the_parser(self):
        """ "this is not an office document" is the parser's error, with its own wire code.
        Pre-empting it here would report a corrupt upload as an oversized one."""
        assert_archive_within_limits(b"not a zip at all")  # must not raise

    def test_the_guard_reads_in_bounded_chunks(self):
        """A guard that has to decompress the whole bomb to decide IS the bomb.

        ``_actual_size`` stops one byte past the cap and reads a chunk at a time, so peak
        memory is one chunk rather than the whole part — measured at 0.04 s for a
        32,768-part archive. Pinned on the source because the alternative (`part.read()`
        with no argument) is a silent one-word regression that no size fixture detects.
        """
        source = Path(limits.__file__).read_text(encoding="utf-8")
        assert "part.read(1 << 20)" in source
        assert "part.read()" not in source


# ── structural caps: truncate, and say what was dropped ─────────────────────


class TestBlockCap:
    @pytest.mark.timeout(600)
    def test_the_issue_repro_is_bounded_at_the_real_cap(self):
        """#2747's document shape against the SHIPPED cap, not a patched one.

        Asserts the block COUNT rather than wall time: the count is what bounds the work
        deterministically, and a timing assertion would flake (this machine runs six suites
        at once). Measured for the record — the issue's 200,000-paragraph 56 KB file took
        153 s on ``main`` and 7.7 s here, on an unloaded host.

        30,000 source paragraphs rather than 200,000 because the walk stops at the cap
        either way, so the extra 170,000 buy nothing but XML generation. The raised timeout
        is for the 20,000-block walk itself, which is the point of the test: the suite's
        120 s default was written for unit tests, and no assertion is relaxed to fit it.
        """
        from personalclaw.documents.docx_parser import parse_docx

        model, loss = parse_docx(_docx_with(30_000))
        assert len(model.blocks) <= MAX_BLOCKS + 1
        assert len(model.blocks) >= MAX_BLOCKS, "truncation must keep everything up to the cap"
        assert loss.of_kind("size_limit"), "a truncated parse that reports nothing is a lie"

    def test_truncation_names_the_way_out(self, monkeypatch):
        """A cap with no escape is the outage. The loss must point at the raw route, which
        serves the whole file untouched."""
        from personalclaw.documents import docx_parser

        monkeypatch.setattr(docx_parser, "MAX_BLOCKS", 3)
        model, loss = docx_parser.parse_docx(_docx_with(50))
        assert len(model.blocks) <= 4
        items = loss.of_kind("size_limit")
        assert len(items) == 1, "the cap must report exactly once, not per dropped block"
        assert "raw" in items[0].detail

    def test_a_document_under_the_cap_reports_nothing(self, monkeypatch):
        from personalclaw.documents import docx_parser

        monkeypatch.setattr(docx_parser, "MAX_BLOCKS", 100)
        model, loss = docx_parser.parse_docx(_docx_with(5))
        assert not loss.of_kind("size_limit")
        assert len(model.blocks) == 6


class TestCellCap:
    def _workbook(self, rows: int, cols: int) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        for row in range(rows):
            ws.append([f"v{row}-{col}" for col in range(cols)])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        return buf.getvalue()

    def test_the_budget_is_the_WORKBOOK_not_the_sheet(self, monkeypatch):
        """A per-sheet cap is evaded by adding sheets, so the allowance is the workbook's.

        Two sheets of 6 cells against a 9-cell budget must truncate, and a per-sheet cap
        would have let both through.
        """
        from openpyxl import Workbook

        from personalclaw.documents import xlsx_parser

        monkeypatch.setattr(xlsx_parser, "MAX_CELLS", 9)
        wb = Workbook()
        for index in range(2):
            ws = wb.create_sheet(f"s{index}")
            for row in range(3):
                ws.append(["a", "b"])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        model, loss = xlsx_parser.parse_xlsx(buf.getvalue())
        assert sum(len(r) for s in model.sheets for r in s.cells) <= 9
        assert loss.of_kind("size_limit")

    def test_truncation_drops_whole_rows_and_locates_itself(self, monkeypatch):
        """A truncated sheet must not be a ragged grid, and the loss must say WHERE in the
        sheet's own terms — the locator this parser already uses (``Sales!C2``)."""
        from personalclaw.documents import xlsx_parser

        monkeypatch.setattr(xlsx_parser, "MAX_CELLS", 7)
        model, loss = xlsx_parser.parse_xlsx(self._workbook(rows=10, cols=3))
        widths = {len(r) for s in model.sheets for r in s.cells}
        assert widths in ({3}, set()), f"ragged rows after truncation: {widths}"
        item = loss.of_kind("size_limit")[0]
        assert item.location, "a sheet loss with no location is unfindable"
        assert "row" in item.detail


def test_size_limit_is_reported_by_every_structural_cap(monkeypatch):
    """ONE loss kind, THREE producers — so it is asserted once across all of them.

    ``size_limit`` is the only kind in ``LOSS_KINDS`` that every parser can emit (the sheet
    and deck kinds are format-specific). A per-format assertion would let two of the three
    stop reporting while the vocabulary-coverage rail in ``test_docx_parser`` still read as
    covered. This is the test that rail points at.
    """
    from openpyxl import Workbook
    from pptx import Presentation

    from personalclaw.documents import docx_parser, pptx_parser, xlsx_parser

    monkeypatch.setattr(docx_parser, "MAX_BLOCKS", 2)
    monkeypatch.setattr(xlsx_parser, "MAX_CELLS", 2)
    monkeypatch.setattr(pptx_parser, "MAX_SLIDES", 1)

    wb = Workbook()
    for row in range(4):
        wb.active.append(["a", "b"])
    sheet_buf = io.BytesIO()
    wb.save(sheet_buf)
    wb.close()

    prs = Presentation()
    for index in range(3):
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = f"s{index}"
    deck_buf = io.BytesIO()
    prs.save(deck_buf)

    for label, (_model, loss) in {
        "docx": docx_parser.parse_docx(_docx_with(20)),
        "xlsx": xlsx_parser.parse_xlsx(sheet_buf.getvalue()),
        "pptx": pptx_parser.parse_pptx(deck_buf.getvalue()),
    }.items():
        items = loss.of_kind("size_limit")
        assert items, f"{label} truncated without reporting it"
        assert "raw" in items[0].detail, f"{label}'s loss does not name the way out"


class TestSlideCap:
    def test_a_deck_over_the_cap_truncates(self, monkeypatch):
        from pptx import Presentation

        from personalclaw.documents import pptx_parser

        monkeypatch.setattr(pptx_parser, "MAX_SLIDES", 2)
        prs = Presentation()
        layout = prs.slide_layouts[1]
        for index in range(5):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = f"Slide {index}"
        buf = io.BytesIO()
        prs.save(buf)
        model, loss = pptx_parser.parse_pptx(buf.getvalue())
        assert len(model.slides) <= 2
        item = loss.of_kind("size_limit")[0]
        assert "5" in item.detail, "the loss must say how many slides the deck declared"


# ── the PROPERTY, over every parser the codebase can discover ───────────────


class TestEveryParserInheritsTheCaps:
    """An enumerated rail ("docx, xlsx and pptx are capped") cannot see the parser added
    tomorrow. These two derive their population from the code instead — one from the codec
    table the wire actually serves, one from the module's own ``parse_*`` functions.

    **What they CANNOT see, stated plainly:** a parser that lives outside
    ``documents/`` and is not registered as a ``ModelCodec``. ``doc_parser.py`` is exactly
    that shape and carries its own (older, separate) caps; a THIRD such module would be
    invisible to both rails.
    """

    def test_every_declared_model_kind_refuses_an_oversized_package(self):
        from personalclaw.documents.model_codec import MODEL_KINDS, get_codec

        assert len(MODEL_KINDS) >= 3, f"only {len(MODEL_KINDS)} kinds — vacuous"
        oversized = _oversized_package()
        for kind in MODEL_KINDS:
            codec = get_codec(kind)
            assert codec is not None
            with pytest.raises(DocumentTooLarge):
                codec.parse(oversized)

    def test_every_parse_entry_point_calls_the_archive_check(self):
        """AST, so a fourth ``parse_*`` added to ``documents/`` is covered on the day it is
        written — including one no codec has registered yet."""
        package = Path(limits.__file__).parent
        found: dict[str, bool] = {}
        for py in sorted(package.glob("*_parser.py")):
            tree = ast.parse(py.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if not isinstance(node, ast.FunctionDef) or not node.name.startswith("parse_"):
                    continue
                calls = {
                    child.func.id
                    for child in ast.walk(node)
                    if isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
                }
                found[f"{py.name}::{node.name}"] = "assert_archive_within_limits" in calls
        assert len(found) >= 3, f"AST scan found only {sorted(found)} — vacuous"
        missing = sorted(name for name, ok in found.items() if not ok)
        assert not missing, f"parse entry points with no archive cap: {missing}"


# ── calibration: the caps must not refuse ordinary documents ────────────────


class TestKnownCleanCorpus:
    """The floor that makes the caps a control rather than an outage.

    Calibrated against every real ``.docx``/``.xlsx``/``.pptx`` on the development machine
    (Word-authored specifications, Excel-authored financial models) plus documents at the
    top of the ORDINARY range — because the real corpus alone would have picked caps far
    too tight. Only the machine-independent half can live in a test, so that is what this
    asserts.
    """

    def test_the_repo_fixture_parses_untruncated(self):
        from personalclaw.documents.docx_parser import parse_docx

        model, loss = parse_docx((_FIXTURES / "word_authored.docx").read_bytes())
        assert not loss.of_kind("size_limit")
        assert model.blocks

    def test_a_library_template_parses_untruncated(self):
        """python-docx's own default template is 908 KB of XML — the largest single XML
        total in the measured corpus, and comfortably inside the caps."""
        import docx

        from personalclaw.documents.docx_parser import parse_docx

        template = Path(docx.__file__).parent / "templates" / "default.docx"
        model, loss = parse_docx(template.read_bytes())
        assert not loss.of_kind("size_limit")
        assert model is not None

    def test_an_ordinary_large_spreadsheet_is_not_refused(self):
        """THE calibration regression test.

        A 2,000 x 100 export — 200,000 cells, a normal thing to be handed — carries
        **10.9 MB of XML in one part**. A per-part cap chosen from the largest real file
        (438 KB) with "generous" headroom (8 MB) refuses it. This test is what makes that
        mistake red instead of a support ticket.
        """
        from openpyxl import Workbook

        from personalclaw.documents.xlsx_parser import parse_xlsx

        wb = Workbook()
        ws = wb.active
        for row in range(2000):
            ws.append([f"v{row}-{col}" for col in range(100)])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        data = buf.getvalue()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            biggest = max(info.file_size for info in archive.infolist())
        assert biggest > 8 * 1024 * 1024, (
            "this workbook no longer exceeds 8 MB in one part, so it no longer guards the "
            f"calibration it was written for (largest part {biggest})"
        )
        model, loss = parse_xlsx(data)
        assert not loss.of_kind("size_limit")
        assert sum(len(r) for s in model.sheets for r in s.cells) == 200_000

    def test_the_caps_leave_real_headroom(self):
        """Pinned so a later "tighten it a bit" has to argue with the corpus.

        Largest values measured across the corpus: 251 blocks, 7,280 cells, 12 slides,
        43 archive parts, 438 KB in one XML part.
        """
        assert MAX_BLOCKS >= 251 * 20
        assert MAX_CELLS >= 7_280 * 20
        assert MAX_SLIDES >= 12 * 20
        assert MAX_ARCHIVE_ENTRIES >= 2_046 * 8  # a 1,000-slide deck is 2,046 parts
        assert MAX_XML_PART_BYTES >= 33 * 1024 * 1024  # a 600k-cell sheet is 33 MB
